"""Gerador didático de produtos topográficos a partir de CSV.

Uso recomendado:
    python gerador_planta.py --input PONTOS.csv --output saida --source-crs EPSG:31985

O CRS EPSG:31985 é somente um exemplo para os dados didáticos. Em um projeto
real, substitua-o pelo CRS efetivamente informado no levantamento.
"""

from __future__ import annotations

import argparse
import csv
import math
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

REQUIRED_FIELDS = ["ID", "X", "Y", "Z", "CODIGO", "DESCRICAO"]
LAYER_BY_CODE = {
    "PV": "TOPO_POSTE",
    "ED": "TOPO_EDIFICACAO",
}
KML_NS = "http://www.opengis.net/kml/2.2"
ET.register_namespace("", KML_NS)


@dataclass(frozen=True)
class PointRecord:
    point_id: int
    x: float
    y: float
    z: float
    code: str
    description: str
    source_line: int


def as_float(value: str, field: str, line_number: int) -> float:
    """Converte um campo numérico e informa linha/campo em caso de erro."""
    text = value.strip()
    if not text:
        raise ValueError(f"linha {line_number}: campo {field} vazio")
    # Permite decimal com vírgula apenas quando não há ponto decimal.
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError as exc:
        raise ValueError(f"linha {line_number}: campo {field} não numérico: {value!r}") from exc
    if not math.isfinite(number):
        raise ValueError(f"linha {line_number}: campo {field} não finito")
    return number


def detect_dialect(sample: str) -> csv.Dialect:
    """Detecta delimitador comum, mantendo uma lista controlada de opções."""
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        class DefaultDialect(csv.excel):
            delimiter = ","
        return DefaultDialect()


def read_points(path: Path) -> tuple[list[PointRecord], list[str], str]:
    """Lê e valida registros. Registros inválidos não entram nos produtos."""
    raw = path.read_text(encoding="utf-8-sig")
    dialect = detect_dialect(raw[:4096])
    records: list[PointRecord] = []
    issues: list[str] = []
    seen_ids: dict[int, int] = {}
    seen_xy: dict[tuple[float, float], int] = {}

    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle, dialect=dialect)
        fieldnames = reader.fieldnames or []
        missing = [field for field in REQUIRED_FIELDS if field not in fieldnames]
        if missing:
            raise ValueError(f"campos obrigatórios ausentes: {', '.join(missing)}")

        for line_number, row in enumerate(reader, start=2):
            try:
                point_id_value = row.get("ID", "").strip()
                if not point_id_value:
                    raise ValueError(f"linha {line_number}: campo ID vazio")
                point_id_float = as_float(point_id_value, "ID", line_number)
                if not point_id_float.is_integer():
                    raise ValueError(f"linha {line_number}: ID não inteiro")
                point_id = int(point_id_float)
                x = as_float(row.get("X", ""), "X", line_number)
                y = as_float(row.get("Y", ""), "Y", line_number)
                z = as_float(row.get("Z", ""), "Z", line_number)
                code = row.get("CODIGO", "").strip().upper()
                description = row.get("DESCRICAO", "").strip()
                if not code:
                    raise ValueError(f"linha {line_number}: campo CODIGO vazio")
                if not description:
                    raise ValueError(f"linha {line_number}: campo DESCRICAO vazio")
                if point_id in seen_ids:
                    raise ValueError(
                        f"linha {line_number}: ID duplicado; primeira ocorrência na linha {seen_ids[point_id]}"
                    )
                xy_key = (round(x, 6), round(y, 6))
                if xy_key in seen_xy:
                    issues.append(
                        f"ALERTA linha {line_number}: coordenada X/Y repetida; primeira ocorrência na linha {seen_xy[xy_key]}"
                    )
                seen_ids[point_id] = line_number
                seen_xy[xy_key] = line_number
                if code not in LAYER_BY_CODE:
                    issues.append(
                        f"ALERTA linha {line_number}: código {code!r} não mapeado; enviado para TOPO_REVISAR"
                    )
                records.append(PointRecord(point_id, x, y, z, code, description, line_number))
            except ValueError as exc:
                issues.append(f"ERRO {exc}")

    return records, issues, dialect.delimiter


def layer_for(record: PointRecord) -> str:
    return LAYER_BY_CODE.get(record.code, "TOPO_REVISAR")


def distance_azimuth(first: PointRecord, second: PointRecord) -> tuple[float, float]:
    """Distância e azimute em plano, contado do Norte para Leste."""
    dx = second.x - first.x
    dy = second.y - first.y
    distance = math.hypot(dx, dy)
    azimuth = math.degrees(math.atan2(dx, dy)) % 360.0
    return distance, azimuth


def ensure_output(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def write_dxf(records: Iterable[PointRecord], path: Path, connect: bool = False) -> None:
    """Escreve pontos, textos e quadro em DXF usando ezdxf."""
    try:
        import ezdxf
    except ImportError as exc:
        raise RuntimeError("instale ezdxf para gerar DXF: pip install ezdxf") from exc

    ensure_output(path)
    record_list = list(records)
    document = ezdxf.new("R2010")
    modelspace = document.modelspace()
    layers = {
        "TOPO_PONTOS": 7,
        "TOPO_ID": 2,
        "TOPO_COTAS": 3,
        "TOPO_DESC": 4,
        "TOPO_POSTE": 1,
        "TOPO_EDIFICACAO": 5,
        "TOPO_REVISAR": 6,
        "TOPO_QUADRO": 7,
    }
    for name, color in layers.items():
        if name not in document.layers:
            document.layers.add(name=name, color=color)

    for record in record_list:
        modelspace.add_point((record.x, record.y, record.z), dxfattribs={"layer": "TOPO_PONTOS"})
        modelspace.add_text(
            str(record.point_id),
            dxfattribs={"layer": "TOPO_ID", "height": 1.0},
        ).set_placement((record.x + 1.0, record.y + 1.0, record.z))
        modelspace.add_text(
            f"{record.z:.2f}",
            dxfattribs={"layer": "TOPO_COTAS", "height": 1.0},
        ).set_placement((record.x + 1.0, record.y, record.z))
        modelspace.add_text(
            record.description,
            dxfattribs={"layer": "TOPO_DESC", "height": 1.0},
        ).set_placement((record.x + 1.0, record.y - 1.0, record.z))
        modelspace.add_text(
            record.code,
            dxfattribs={"layer": layer_for(record), "height": 0.8},
        ).set_placement((record.x + 2.0, record.y + 2.0, record.z))

    if record_list:
        base_x = max(record.x for record in record_list) + 20.0
        base_y = max(record.y for record in record_list)
        modelspace.add_text(
            "ID | X | Y | Z | CODIGO",
            dxfattribs={"layer": "TOPO_QUADRO", "height": 1.0},
        ).set_placement((base_x, base_y, 0.0))
        for row_index, record in enumerate(record_list, start=1):
            table_text = f"{record.point_id} | {record.x:.3f} | {record.y:.3f} | {record.z:.2f} | {record.code}"
            modelspace.add_text(
                table_text,
                dxfattribs={"layer": "TOPO_QUADRO", "height": 1.0},
            ).set_placement((base_x, base_y - 2.0 * row_index, 0.0))

    if connect and len(record_list) >= 2:
        # Somente para demonstração: a conectividade precisa ser uma regra do projeto.
        modelspace.add_lwpolyline(
            [(record.x, record.y) for record in record_list],
            dxfattribs={"layer": "TOPO_QUADRO"},
        )

    document.saveas(path)


def write_kml(records: Iterable[PointRecord], path: Path, source_crs: str) -> None:
    """Transforma pontos para EPSG:4326 e escreve KML."""
    try:
        from pyproj import Transformer
    except ImportError as exc:
        raise RuntimeError("instale pyproj para gerar KML: pip install pyproj") from exc

    transformer = Transformer.from_crs(source_crs, "EPSG:4326", always_xy=True)
    ensure_output(path)
    kml = ET.Element(f"{{{KML_NS}}}kml")
    document = ET.SubElement(kml, f"{{{KML_NS}}}Document")
    ET.SubElement(document, f"{{{KML_NS}}}name").text = "Pontos topográficos"
    for record in records:
        longitude, latitude = transformer.transform(record.x, record.y)
        placemark = ET.SubElement(document, f"{{{KML_NS}}}Placemark")
        ET.SubElement(placemark, f"{{{KML_NS}}}name").text = f"P-{record.point_id}"
        ET.SubElement(placemark, f"{{{KML_NS}}}description").text = (
            f"{record.code} - {record.description}; cota={record.z:.2f}"
        )
        point = ET.SubElement(placemark, f"{{{KML_NS}}}Point")
        ET.SubElement(point, f"{{{KML_NS}}}coordinates").text = (
            f"{longitude:.9f},{latitude:.9f},{record.z:.3f}"
        )
    if hasattr(ET, "indent"):
        ET.indent(kml, space="  ")
    ET.ElementTree(kml).write(path, encoding="utf-8", xml_declaration=True)


def write_xlsx(records: Iterable[PointRecord], path: Path, source_crs: str) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("instale openpyxl para gerar XLSX: pip install openpyxl") from exc

    ensure_output(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Coordenadas"
    headers = ["ID", "X", "Y", "Z", "CODIGO", "DESCRICAO", "LAYER", "CRS"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for record in records:
        sheet.append([
            record.point_id,
            record.x,
            record.y,
            record.z,
            record.code,
            record.description,
            layer_for(record),
            source_crs,
        ])
    for column, width in {"A": 10, "B": 16, "C": 16, "D": 12, "E": 12, "F": 24, "G": 20, "H": 16}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)


def write_docx(records: Iterable[PointRecord], issues: list[str], path: Path, source_crs: str) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("instale python-docx para gerar DOCX: pip install python-docx") from exc

    ensure_output(path)
    record_list = list(records)
    document = Document()
    document.add_heading("Relatório didático de automação topográfica", level=1)
    document.add_paragraph(
        "Este documento é um exemplo pedagógico gerado a partir de uma base CSV validada. "
        "Ele não substitui relatório técnico, memorial descritivo ou responsabilidade profissional."
    )
    document.add_heading("Parâmetros", level=2)
    document.add_paragraph(f"CRS de origem informado: {source_crs}")
    document.add_paragraph(f"Registros aprovados: {len(record_list)}")
    document.add_paragraph(f"Data UTC da execução: {datetime.now(timezone.utc).isoformat()}")
    document.add_heading("Coordenadas", level=2)
    table = document.add_table(rows=1, cols=6)
    for cell, text in zip(table.rows[0].cells, ["ID", "X", "Y", "Z", "Código", "Descrição"]):
        cell.text = text
    for record in record_list:
        cells = table.add_row().cells
        values = [record.point_id, f"{record.x:.3f}", f"{record.y:.3f}", f"{record.z:.2f}", record.code, record.description]
        for cell, value in zip(cells, values):
            cell.text = str(value)
    document.add_heading("Ocorrências de validação", level=2)
    if issues:
        for issue in issues:
            document.add_paragraph(issue, style="List Bullet")
    else:
        document.add_paragraph("Nenhuma ocorrência registrada.")
    document.save(path)


def write_validation_report(records: list[PointRecord], issues: list[str], delimiter: str, path: Path, source_crs: str) -> None:
    ensure_output(path)
    lines = [
        "RELATÓRIO DE VALIDAÇÃO — GERADOR DIDÁTICO",
        f"Execução UTC: {datetime.now(timezone.utc).isoformat()}",
        f"Registros aprovados: {len(records)}",
        f"Delimitador detectado: {delimiter!r}",
        f"CRS informado pelo operador: {source_crs}",
        "",
        "OCORRÊNCIAS:",
    ]
    lines.extend(issues or ["Nenhuma ocorrência registrada."])
    lines.extend([
        "",
        "NOTA: CRS e unidades não são inferidos automaticamente neste material.",
        "NOTA: conectividade de linhas deve ser definida por regra do projeto.",
    ])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="CSV de pontos")
    parser.add_argument("--output", type=Path, default=Path("saida"), help="diretório de saída")
    parser.add_argument("--source-crs", required=True, help="CRS de origem, por exemplo EPSG:31985")
    parser.add_argument("--connect", action="store_true", help="gera polilinha didática na ordem do arquivo")
    args = parser.parse_args()

    records, issues, delimiter = read_points(args.input)
    args.output.mkdir(parents=True, exist_ok=True)
    write_dxf(records, args.output / "planta.dxf", connect=args.connect)
    write_kml(records, args.output / "pontos.kml", args.source_crs)
    write_xlsx(records, args.output / "coordenadas.xlsx", args.source_crs)
    write_docx(records, issues, args.output / "memorial.docx", args.source_crs)
    write_validation_report(records, issues, delimiter, args.output / "relatorio_validacao.txt", args.source_crs)
    print(f"Concluído: {len(records)} registros aprovados; {len(issues)} ocorrência(s).")
    print(f"Saídas: {args.output.resolve()}")


if __name__ == "__main__":
    main()
