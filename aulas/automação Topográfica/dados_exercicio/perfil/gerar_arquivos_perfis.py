from pathlib import Path
import csv
import math
from statistics import mean, pstdev
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList

ROOT = Path('/home/ubuntu/projeto_lev_topografico/aulas/automação Topográfica/dados_exercicio/perfil')
ROOT.mkdir(parents=True, exist_ok=True)
SOURCE = Path('/home/ubuntu/upload/pontos.txt')

rows = []
for line in SOURCE.read_text(encoding='utf-8').splitlines():
    if line.strip():
        pid, x, y, z = line.split(',')
        rows.append({'id': int(pid), 'x': float(x), 'y': float(y), 'z': float(z)})

axis = rows[:87]
collected_local = rows[87:]

# The supplied file contains a global-coordinate axis followed by a local-coordinate
# collected set. The anchor pair is point 1 (axis) and point 151 (collected).
dx = axis[0]['x'] - next(r for r in collected_local if r['id'] == 151)['x']
dy = axis[0]['y'] - next(r for r in collected_local if r['id'] == 151)['y']

for r in axis:
    r['station_m'] = 0.0
for previous, current in zip(axis, axis[1:]):
    current['station_m'] = previous['station_m'] + math.hypot(current['x'] - previous['x'], current['y'] - previous['y'])

for r in collected_local:
    r['x_global'] = r['x'] + dx
    r['y_global'] = r['y'] + dy
    nearest = min(axis, key=lambda p: math.hypot(p['x'] - r['x_global'], p['y'] - r['y_global']))
    r['match_id'] = nearest['id']
    r['xy_residual_m'] = math.hypot(nearest['x'] - r['x_global'], nearest['y'] - r['y_global'])
    r['z_provided'] = nearest['z']
    r['dz_m'] = r['z'] - nearest['z']


def azimuth(a, b):
    return math.degrees(math.atan2(b['x'] - a['x'], b['y'] - a['y'])) % 360.0


def interpolate_at(station):
    if station <= 0:
        a, b = axis[0], axis[1]
    elif station >= axis[-1]['station_m']:
        a, b = axis[-2], axis[-1]
    else:
        for a, b in zip(axis, axis[1:]):
            if a['station_m'] <= station <= b['station_m']:
                break
    seg_len = math.hypot(b['x'] - a['x'], b['y'] - a['y'])
    ratio = 0.0 if seg_len == 0 else (station - a['station_m']) / seg_len
    ratio = max(0.0, min(1.0, ratio))
    return {
        'station_m': station,
        'x': a['x'] + ratio * (b['x'] - a['x']),
        'y': a['y'] + ratio * (b['y'] - a['y']),
        'z': a['z'] + ratio * (b['z'] - a['z']),
        'azimuth_deg': azimuth(a, b),
    }


stationing = [interpolate_at(float(s)) for s in range(0, int(axis[-1]['station_m']) // 20 * 20 + 1, 20)]
if stationing[-1]['station_m'] < axis[-1]['station_m'] - 1e-9:
    stationing.append(interpolate_at(axis[-1]['station_m']))

sections = []
for station in [0.0, 500.0, 1000.0, 1500.0]:
    center = interpolate_at(station)
    angle = math.radians(center['azimuth_deg'])
    # Positive offset is the right-hand side of the forward tangent; negative is left.
    for offset in [-10.0, 0.0, 10.0]:
        x = center['x'] + offset * math.cos(angle)
        y = center['y'] - offset * math.sin(angle)
        sections.append({
            'station_m': station,
            'x': x,
            'y': y,
            'z_axis': center['z'],
            'azimuth_deg': center['azimuth_deg'],
            'offset_m': offset,
            'side': 'Eixo' if offset == 0 else ('Esquerda' if offset < 0 else 'Direita'),
            'z_pe3d': None,
            'difference_m': None,
        })

perimeter = []
for a, b in zip(axis, axis[1:]):
    dist = math.hypot(b['x'] - a['x'], b['y'] - a['y'])
    dz = b['z'] - a['z']
    perimeter.append({
        'from_id': a['id'], 'to_id': b['id'], 'station_from_m': a['station_m'],
        'azimuth_deg': azimuth(a, b), 'distance_m': dist, 'z_from_m': a['z'],
        'z_to_m': b['z'], 'dz_m': dz, 'slope_pct': (dz / dist * 100.0) if dist else None,
    })


def write_csv(path, headers, records):
    with path.open('w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(records)

write_csv(ROOT / 'pontos_eixo_fornecidos.csv', ['ID', 'X_m', 'Y_m', 'Z_m', 'Estacao_m'], [
    {'ID': p['id'], 'X_m': f"{p['x']:.4f}", 'Y_m': f"{p['y']:.4f}", 'Z_m': f"{p['z']:.4f}", 'Estacao_m': f"{p['station_m']:.4f}"} for p in axis
])
write_csv(ROOT / 'pontos_coletados_pe3d.csv', ['ID_COLETADO', 'X_LOCAL_m', 'Y_LOCAL_m', 'Z_PE3D_m', 'X_GLOBAL_ESTIMADO_m', 'Y_GLOBAL_ESTIMADO_m', 'ID_FORNECIDO_MATCH', 'RESIDUO_XY_m', 'Z_FORNECIDA_m', 'DIFERENCA_Z_m'], [
    {'ID_COLETADO': p['id'], 'X_LOCAL_m': f"{p['x']:.4f}", 'Y_LOCAL_m': f"{p['y']:.4f}", 'Z_PE3D_m': f"{p['z']:.4f}", 'X_GLOBAL_ESTIMADO_m': f"{p['x_global']:.4f}", 'Y_GLOBAL_ESTIMADO_m': f"{p['y_global']:.4f}", 'ID_FORNECIDO_MATCH': p['match_id'], 'RESIDUO_XY_m': f"{p['xy_residual_m']:.6f}", 'Z_FORNECIDA_m': f"{p['z_provided']:.4f}", 'DIFERENCA_Z_m': f"{p['dz_m']:.4f}"} for p in collected_local
])
write_csv(ROOT / 'estaqueamento_20m.csv', ['Estaca', 'Estacao_m', 'X_m', 'Y_m', 'Cota_interpolada_m', 'Azimute_eixo_deg'], [
    {'Estaca': f"{int(p['station_m'] // 1000)}+{p['station_m'] % 1000:06.2f}", 'Estacao_m': f"{p['station_m']:.4f}", 'X_m': f"{p['x']:.4f}", 'Y_m': f"{p['y']:.4f}", 'Cota_interpolada_m': f"{p['z']:.4f}", 'Azimute_eixo_deg': f"{p['azimuth_deg']:.6f}"} for p in stationing
])
write_csv(ROOT / 'secoes_transversais_500m_10m.csv', ['Estacao_m', 'Estaca', 'Offset_m', 'Lado', 'X_PE3D_alvo_m', 'Y_PE3D_alvo_m', 'Cota_eixo_m', 'Cota_PE3D_extraida_m', 'Diferenca_PE3D_eixo_m'], [
    {'Estacao_m': f"{p['station_m']:.2f}", 'Estaca': f"{int(p['station_m'] // 1000)}+{p['station_m'] % 1000:06.2f}", 'Offset_m': f"{p['offset_m']:.2f}", 'Lado': p['side'], 'X_PE3D_alvo_m': f"{p['x']:.4f}", 'Y_PE3D_alvo_m': f"{p['y']:.4f}", 'Cota_eixo_m': f"{p['z_axis']:.4f}", 'Cota_PE3D_extraida_m': '', 'Diferenca_PE3D_eixo_m': ''} for p in sections
])
write_csv(ROOT / 'roteiro_perimetrico_eixo.csv', ['De', 'Para', 'Estacao_de_m', 'Azimute_deg', 'Distancia_m', 'Cota_de_m', 'Cota_para_m', 'Delta_Z_m', 'Declividade_pct'], [
    {'De': p['from_id'], 'Para': p['to_id'], 'Estacao_de_m': f"{p['station_from_m']:.4f}", 'Azimute_deg': f"{p['azimuth_deg']:.6f}", 'Distancia_m': f"{p['distance_m']:.4f}", 'Cota_de_m': f"{p['z_from_m']:.4f}", 'Cota_para_m': f"{p['z_to_m']:.4f}", 'Delta_Z_m': f"{p['dz_m']:.4f}", 'Declividade_pct': f"{p['slope_pct']:.4f}"} for p in perimeter
])
write_csv(ROOT / 'comparacao_cotas_fornecida_pe3d.csv', ['ID_COLETADO', 'ID_FORNECIDO', 'Cota_fornecida_m', 'Cota_coletada_PE3D_m', 'Diferenca_PE3D_m', 'Residuo_XY_m', 'Observacao'], [
    {'ID_COLETADO': p['id'], 'ID_FORNECIDO': p['match_id'], 'Cota_fornecida_m': f"{p['z_provided']:.4f}", 'Cota_coletada_PE3D_m': f"{p['z']:.4f}", 'Diferenca_PE3D_m': f"{p['dz_m']:.4f}", 'Residuo_XY_m': f"{p['xy_residual_m']:.6f}", 'Observacao': 'Pareamento inicial por translação; confirmar no QGIS.'} for p in collected_local
])

# Workbook styling
wb = Workbook()
ws = wb.active
ws.title = 'Visao_geral'
navy = '0B1F3A'; teal = '007C83'; gold = 'D6A944'; light = 'E2E8F0'; paper = 'F5F7FA'; muted = '475569'; red = 'FCE4D6'
thin = Side(style='thin', color='B8C2CC')
medium = Side(style='medium', color=navy)

def setup_sheet(sheet, widths=None):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'B5'
    sheet.column_dimensions['A'].width = 3
    if widths:
        for col, width in widths.items():
            sheet.column_dimensions[col].width = width

def title(sheet, text, subtitle=None, end_col=10):
    sheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=end_col)
    cell = sheet.cell(2, 2, text)
    cell.font = Font(name='Aptos Display', size=18, bold=True, color=navy)
    cell.fill = PatternFill('solid', fgColor=paper)
    cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[2].height = 28
    if subtitle:
        sheet.merge_cells(start_row=3, start_column=2, end_row=3, end_column=end_col)
        c = sheet.cell(3, 2, subtitle)
        c.font = Font(name='Aptos', size=10, italic=True, color=muted)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[3].height = 30

def header_row(sheet, row, headers):
    for col, value in enumerate(headers, start=2):
        c = sheet.cell(row, col, value)
        c.font = Font(name='Aptos', size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=teal)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = Border(top=medium, bottom=medium, left=thin, right=thin)
    sheet.row_dimensions[row].height = 30

def style_table(sheet, start_row, end_row, start_col, end_col, number_formats=None):
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            c = sheet.cell(row, col)
            c.font = Font(name='Aptos', size=10, color=navy)
            c.fill = PatternFill('solid', fgColor='FFFFFF' if row % 2 else 'F8FAFC')
            c.border = Border(bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical='center', wrap_text=False)
            if number_formats and col in number_formats:
                c.number_format = number_formats[col]
    sheet.auto_filter.ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"

def link_cell(sheet, cell_ref, text, target):
    c = sheet[cell_ref]
    c.value = text
    c.hyperlink = target
    c.font = Font(name='Aptos', size=10, color=teal, underline='single')

setup_sheet(ws, {'B': 32, 'C': 74, 'D': 22, 'E': 22, 'F': 22, 'G': 22, 'H': 22})
title(ws, 'Exercício de Perfis Topográficos — Estaqueamento e PE3D', 'Base: pontos.txt fornecido pelo professor. A planilha organiza o eixo, o roteiro perimétrico, o estaqueamento, as seções transversais e a comparação de cotas.', 8)
ws['B5'] = 'OBJETIVO DO ARQUIVO'; ws['B5'].font = Font(size=12, bold=True, color=gold)
ws.merge_cells('B6:H7'); ws['B6'] = 'Usar os pontos fornecidos para gerar, no CAD e no QGIS, um perfil longitudinal com estaqueamento de 20 em 20 metros e seções transversais a cada 500 m, com 10 m para cada lado. As cotas das seções transversais devem ser extraídas do PE3D e preenchidas na aba correspondente.'; ws['B6'].alignment = Alignment(wrap_text=True, vertical='top'); ws['B6'].font = Font(size=11, color=navy)
ws['B9'] = 'RESUMO DA BASE'; ws['B9'].font = Font(size=12, bold=True, color=gold)
summary = [
    ('Pontos do eixo fornecidos', len(axis)),
    ('Extensão acumulada do eixo (m)', axis[-1]['station_m']),
    ('Pontos coletados/PE3D no arquivo', len(collected_local)),
    ('Translação inicial X (m)', dx),
    ('Translação inicial Y (m)', dy),
    ('Resíduo XY máximo do pareamento (m)', max(p['xy_residual_m'] for p in collected_local)),
    ('Diferença Z média PE3D - fornecida (m)', mean(p['dz_m'] for p in collected_local)),
]
for idx, (label, value) in enumerate(summary, start=10):
    ws.cell(idx, 2, label).font = Font(bold=True, color=navy)
    ws.cell(idx, 3, value).font = Font(color=navy)
    ws.cell(idx, 3).number_format = '#,##0.0000'
    ws.cell(idx, 2).fill = PatternFill('solid', fgColor=light if idx % 2 == 0 else 'FFFFFF')
    ws.cell(idx, 3).fill = PatternFill('solid', fgColor=light if idx % 2 == 0 else 'FFFFFF')
ws['B19'] = 'ABAS'; ws['B19'].font = Font(size=12, bold=True, color=gold)
links = [('Dados_Eixo', 'Eixo fornecido e estação acumulada'), ('Pontos_PE3D', 'Pontos coletados/PE3D e pareamento inicial'), ('Roteiro_Perimetrico', 'Azimute, distância, cotas e declividade'), ('Estaqueamento_20m', 'Estações a cada 20 m e cotas interpoladas'), ('Secoes_500m', 'Alvos de ±10 m para extração no PE3D'), ('Comparacao_Cotas', 'Diferença de cota entre conjuntos')]
for i, (sheet_name, desc) in enumerate(links, start=20):
    link_cell(ws, f'B{i}', sheet_name, f"#'{sheet_name}'!A1")
    ws.cell(i, 3, desc).font = Font(color=muted)
ws['B29'] = 'Premissa crítica'; ws['B29'].font = Font(bold=True, color='FFFFFF'); ws['B29'].fill = PatternFill('solid', fgColor=gold)
ws.merge_cells('B30:H31'); ws['B30'] = 'Os registros 1–87 formam o eixo em coordenadas grandes. Os registros 88–166 estão em coordenadas locais e foram pareados preliminarmente por uma translação ancorada no ponto 1 ↔ ponto 151. O grupo deve confirmar essa hipótese no QGIS antes de aceitar a comparação de cotas.'; ws['B30'].alignment = Alignment(wrap_text=True, vertical='top'); ws['B30'].fill = PatternFill('solid', fgColor='FFF8E7')

# Data sheets
ws = wb.create_sheet('Dados_Eixo'); setup_sheet(ws, {'B': 10, 'C': 16, 'D': 18, 'E': 14, 'F': 16, 'G': 18})
title(ws, 'Dados do eixo fornecido', 'Sequência dos pontos 1–87. A estação acumulada é calculada pela distância plana entre pontos consecutivos.', 7)
headers = ['ID', 'X (m)', 'Y (m)', 'Z (m)', 'Estação acumulada (m)', 'Fonte']
header_row(ws, 5, headers)
for r, p in enumerate(axis, start=6):
    vals = [p['id'], p['x'], p['y'], p['z'], p['station_m'], 'pontos.txt — eixo fornecido']
    for c, v in enumerate(vals, start=2): ws.cell(r, c, v)
style_table(ws, 6, 5 + len(axis), 2, 7, {2: '#,##0', 3: '#,##0.0000', 4: '#,##0.0000', 5: '#,##0.0000', 6: '#,##0.0000'})

ws = wb.create_sheet('Pontos_PE3D'); setup_sheet(ws, {'B': 14, 'C': 16, 'D': 16, 'E': 16, 'F': 20, 'G': 20, 'H': 16, 'I': 16, 'J': 18, 'K': 18, 'L': 52})
title(ws, 'Pontos coletados / PE3D', 'Coordenadas locais do arquivo original, translação estimada para o sistema do eixo, correspondência espacial preliminar e diferença de cota.', 12)
headers = ['ID coletado', 'X local (m)', 'Y local (m)', 'Z PE3D/coletada (m)', 'X global estimado (m)', 'Y global estimado (m)', 'ID fornecido pareado', 'Resíduo XY (m)', 'Z fornecida (m)', 'Diferença Z (m)', 'Observação']
header_row(ws, 5, headers)
for r, p in enumerate(collected_local, start=6):
    vals = [p['id'], p['x'], p['y'], p['z'], p['x_global'], p['y_global'], p['match_id'], p['xy_residual_m'], p['z_provided'], p['dz_m'], 'Pareamento inicial por translação; confirmar no QGIS.']
    for c, v in enumerate(vals, start=2): ws.cell(r, c, v)
style_table(ws, 6, 5 + len(collected_local), 2, 12, {2: '#,##0', 3: '#,##0.0000', 4: '#,##0.0000', 5: '#,##0.0000', 6: '#,##0.0000', 7: '#,##0.0000', 8: '#,##0', 9: '#,##0.000000', 10: '#,##0.0000', 11: '#,##0.0000'})
for row in range(6, 6 + len(collected_local)):
    if abs(ws.cell(row, 11).value) > 0.0004:
        ws.cell(row, 11).fill = PatternFill('solid', fgColor='FFF2CC')

ws = wb.create_sheet('Roteiro_Perimetrico'); setup_sheet(ws, {'B': 9, 'C': 9, 'D': 18, 'E': 16, 'F': 16, 'G': 15, 'H': 15, 'I': 15, 'J': 16})
title(ws, 'Roteiro perimétrico / longitudinal entre pontos', 'Azimute contado a partir do Norte, no sentido horário. A sequência 1→87 não foi fechada automaticamente; o grupo deve decidir se o conjunto representa eixo ou perímetro.', 10)
headers = ['De', 'Para', 'Estação de (m)', 'Azimute (°)', 'Distância (m)', 'Cota de (m)', 'Cota para (m)', 'ΔZ (m)', 'Declividade (%)']
header_row(ws, 5, headers)
for r, p in enumerate(perimeter, start=6):
    vals = [p['from_id'], p['to_id'], p['station_from_m'], p['azimuth_deg'], p['distance_m'], p['z_from_m'], p['z_to_m'], p['dz_m'], p['slope_pct']]
    for c, v in enumerate(vals, start=2): ws.cell(r, c, v)
style_table(ws, 6, 5 + len(perimeter), 2, 10, {2: '#,##0', 3: '#,##0', 4: '#,##0.0000', 5: '#,##0.000000', 6: '#,##0.0000', 7: '#,##0.0000', 8: '#,##0.0000', 9: '#,##0.0000', 10: '#,##0.0000'})

ws = wb.create_sheet('Estaqueamento_20m'); setup_sheet(ws, {'B': 14, 'C': 16, 'D': 18, 'E': 18, 'F': 22, 'G': 20})
title(ws, 'Estaqueamento a cada 20 metros', 'Coordenadas e cota interpolada linearmente ao longo do eixo. No CAD e no QGIS, o grupo deverá gerar o perfil longitudinal a partir desta tabela e conferir a interpolação.', 7)
headers = ['Estaca', 'Estação (m)', 'X (m)', 'Y (m)', 'Cota interpolada (m)', 'Azimute do eixo (°)']
header_row(ws, 5, headers)
for r, p in enumerate(stationing, start=6):
    estaca = f"{int(p['station_m'] // 1000)}+{p['station_m'] % 1000:06.2f}"
    vals = [estaca, p['station_m'], p['x'], p['y'], p['z'], p['azimuth_deg']]
    for c, v in enumerate(vals, start=2): ws.cell(r, c, v)
style_table(ws, 6, 5 + len(stationing), 2, 7, {3: '#,##0.0000', 4: '#,##0.0000', 5: '#,##0.0000', 6: '#,##0.0000', 7: '#,##0.000000'})
chart = LineChart(); chart.title = 'Perfil longitudinal preliminar'; chart.y_axis.title = 'Cota (m)'; chart.x_axis.title = 'Estação (m)'; chart.height = 7; chart.width = 15
chart.add_data(Reference(ws, min_col=6, min_row=5, max_row=5 + len(stationing)), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=3, min_row=6, max_row=5 + len(stationing)))
chart.legend = None
ws.add_chart(chart, 'I5')

ws = wb.create_sheet('Secoes_500m'); setup_sheet(ws, {'B': 14, 'C': 14, 'D': 12, 'E': 14, 'F': 18, 'G': 18, 'H': 18, 'I': 22, 'J': 24})
title(ws, 'Seções transversais a cada 500 m', 'Para cada seção, medir 10 m à esquerda e 10 m à direita do eixo usando o PE3D. As colunas de cota PE3D e diferença foram deixadas abertas para preenchimento pelos alunos.', 10)
headers = ['Estação (m)', 'Estaca', 'Offset (m)', 'Lado', 'X alvo PE3D (m)', 'Y alvo PE3D (m)', 'Cota eixo (m)', 'Cota PE3D extraída (m)', 'Diferença PE3D−eixo (m)']
header_row(ws, 5, headers)
for r, p in enumerate(sections, start=6):
    vals = [p['station_m'], f"{int(p['station_m'] // 20)}+{p['station_m'] % 20:05.2f}", p['offset_m'], p['side'], p['x'], p['y'], p['z_axis'], None, None]
    for c, v in enumerate(vals, start=2): ws.cell(r, c, v)
    ws.cell(r, 10, f'=IF(I{r}="","",I{r}-H{r})')
style_table(ws, 6, 5 + len(sections), 2, 10, {2: '#,##0.00', 4: '#,##0.00', 6: '#,##0.0000', 7: '#,##0.0000', 8: '#,##0.0000', 9: '#,##0.0000', 10: '#,##0.0000'})
for row in range(6, 6 + len(sections)):
    for col in [9, 10]:
        ws.cell(row, col).fill = PatternFill('solid', fgColor='FFF2CC')

ws = wb.create_sheet('Comparacao_Cotas'); setup_sheet(ws, {'B': 14, 'C': 14, 'D': 20, 'E': 22, 'F': 22, 'G': 16, 'H': 50})
title(ws, 'Diferença entre cotas fornecidas e cotas coletadas / PE3D', 'Diferença preliminar calculada como Z coletada/PE3D − Z fornecida. O pareamento espacial é uma hipótese inicial e deve ser conferido no QGIS.', 8)
headers = ['ID coletado', 'ID fornecido', 'Cota fornecida (m)', 'Cota coletada/PE3D (m)', 'Diferença Z (m)', 'Resíduo XY (m)', 'Observação']
header_row(ws, 5, headers)
for r, p in enumerate(collected_local, start=6):
    vals = [p['id'], p['match_id'], p['z_provided'], p['z'], p['dz_m'], p['xy_residual_m'], 'Diferença preliminar; confirmar o pareamento e a origem da cota.']
    for c, v in enumerate(vals, start=2): ws.cell(r, c, v)
style_table(ws, 6, 5 + len(collected_local), 2, 8, {2: '#,##0', 3: '#,##0', 4: '#,##0.0000', 5: '#,##0.0000', 6: '#,##0.0000', 7: '#,##0.000000'})
ws['B86'] = 'Resumo'; ws['B86'].font = Font(bold=True, color='FFFFFF'); ws['B86'].fill = PatternFill('solid', fgColor=teal)
ws['C86'] = 'Média ΔZ (m)'; ws['D86'] = mean(p['dz_m'] for p in collected_local); ws['D86'].number_format = '#,##0.0000'
ws['E86'] = 'Desvio-padrão (m)'; ws['F86'] = pstdev(p['dz_m'] for p in collected_local); ws['F86'].number_format = '#,##0.0000'
ws['G86'] = 'Máx. abs. ΔZ (m)'; ws['H86'] = max(abs(p['dz_m']) for p in collected_local); ws['H86'].number_format = '#,##0.0000'
for c in range(2, 9): ws.cell(86, c).fill = PatternFill('solid', fgColor=light); ws.cell(86, c).font = Font(bold=True, color=navy)

# Common workbook properties
for sheet in wb.worksheets:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = 'Exercício de Perfis Topográficos · UFPE / LABAT'
    sheet.oddFooter.right.text = 'Página &[Page] de &[Pages]'
    sheet.sheet_view.zoomScale = 90
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = 'auto'

out = ROOT / 'planilha_exercicio_perfis_topograficos.xlsx'
wb.save(out)
print('Created', out)
print('axis_points', len(axis))
print('axis_length_m', round(axis[-1]['station_m'], 4))
print('collected_points', len(collected_local))
print('max_xy_residual_m', round(max(p['xy_residual_m'] for p in collected_local), 6))
print('mean_dz_m', round(mean(p['dz_m'] for p in collected_local), 6))
print('stations', len(stationing))
print('sections_rows', len(sections))
